# -*- coding: utf-8 -*-
"""抽取层：openpyxl 双加载（公式版 + 缓存值版），输出全保真几何数据。
含值/公式/换行/边框码/填充色/合并区/隐藏行列/图片图表数。
无值但带边框的空单元格也保留——边框闭合是分割线索，不能丢。"""
import datetime as _dt
import os

import openpyxl

from .normalize import sanitize_xlsx

STYLE_ONLY_LIMIT = 1_500_000  # 超过则丢弃"仅有样式"的空单元格，降级保速度


def _iso(v):
    if isinstance(v, _dt.datetime):
        # 零点时刻的 datetime 输出纯日期：与 calamine 快路径形态一致（对拍零差异），
        # 也避免下游问答出现 "2026-08-01T00:00:00" 这种对用户不友好的键
        if v.hour == v.minute == v.second == v.microsecond == 0:
            return v.date().isoformat()
        return v.isoformat()
    if isinstance(v, (_dt.date, _dt.time)):
        return v.isoformat()
    return v


def _border_mask(cell) -> int:
    b = cell.border
    m = 0
    for i, side in enumerate((b.top, b.bottom, b.left, b.right)):
        if side is not None and side.style:
            m |= 1 << i
    return m


def _fill_rgb(cell):
    f = cell.fill
    if f is None or f.patternType != "solid":
        return None
    fg = f.fgColor
    if fg is None:
        return None
    if getattr(fg, "type", None) == "rgb" and isinstance(fg.rgb, str):
        if fg.rgb in ("00000000", "FFFFFFFF"):  # 无色/白不作为信号
            return None
        return fg.rgb
    if getattr(fg, "type", None) == "theme":
        return f"theme{fg.theme}"
    return None


def _has_value(v) -> bool:
    return v is not None and not (isinstance(v, str) and v.strip() == "")


def extract_sheet(ws, wsv=None, cal_values=None) -> dict:
    """wsv：openpyxl 缓存值页（常规路径）；cal_values：calamine 值字典（大文件快路径）。
    二者互斥，cal_values 优先。"""
    max_r, max_c = ws.max_row, ws.max_column
    degrade = max_r * max_c > STYLE_ONLY_LIMIT

    cells = {}
    border_only = 0
    n_formula_no_cache = 0
    if cal_values is not None:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                cv = cal_values.get((cell.row, cell.column))
                has_v = _has_value(v) or _has_value(cv)
                if not has_v:
                    bm = _border_mask(cell)
                    fr = _fill_rgb(cell)
                    if not degrade and (bm or fr):
                        cells[(cell.row, cell.column)] = {
                            "v": None, "cv": None, "f": False, "nl": False,
                            "bd": bm, "fm": fr,
                        }
                        border_only += 1
                    continue
                is_formula = isinstance(v, str) and v.startswith("=")
                if is_formula and not _has_value(cv):
                    n_formula_no_cache += 1
                nf = cell.number_format
                cells[(cell.row, cell.column)] = {
                    "v": _iso(v), "cv": _iso(cv),
                    "f": is_formula,
                    "nl": isinstance(v, str) and "\n" in v,
                    "bd": _border_mask(cell), "fm": _fill_rgb(cell),
                    "nf": nf if nf and nf != "General" else None,
                }
        return _finish_sheet(ws, cells, border_only, n_formula_no_cache, degrade)

    for row, vrow in zip(ws.iter_rows(), (wsv or ws).iter_rows()):
        for cell, vcell in zip(row, vrow):
            v = cell.value
            cv = vcell.value
            has_v = _has_value(v) or _has_value(cv)
            if not has_v:
                bm = _border_mask(cell)
                fr = _fill_rgb(cell)
                if not degrade and (bm or fr):
                    cells[(cell.row, cell.column)] = {
                        "v": None, "cv": None, "f": False, "nl": False,
                        "bd": bm, "fm": fr,
                    }
                    border_only += 1
                continue
            is_formula = isinstance(v, str) and v.startswith("=")
            if is_formula and not _has_value(cv):
                n_formula_no_cache += 1
            nf = cell.number_format
            cells[(cell.row, cell.column)] = {
                "v": _iso(v), "cv": _iso(cv),
                "f": is_formula,
                "nl": isinstance(v, str) and "\n" in v,
                "bd": _border_mask(cell), "fm": _fill_rgb(cell),
                # 数字显示口径：0.85 按 0% 格式实为 85，下游语义层必须知道
                "nf": nf if nf and nf != "General" else None,
            }

    hidden_rows = {r for r, d in ws.row_dimensions.items() if d.hidden}
    from openpyxl.utils import column_index_from_string
    hidden_cols = {column_index_from_string(c) for c, d in ws.column_dimensions.items()
                   if d.hidden}

    return _finish_sheet(ws, cells, border_only, n_formula_no_cache, degrade,
                         max_r=max_r, max_c=max_c)


def _finish_sheet(ws, cells, border_only, n_formula_no_cache, degrade,
                  max_r=None, max_c=None):
    max_r = max_r or ws.max_row
    max_c = max_c or ws.max_column
    return {
        "name": ws.title,
        "declared_dims": [max_r, max_c],
        "merged": [[m.min_row, m.min_col, m.max_row, m.max_col] for m in ws.merged_cells.ranges],
        "hidden_rows": sorted({r for r, d in ws.row_dimensions.items() if d.hidden}),
        "hidden_cols": sorted(_col_idx(c) for c, d in ws.column_dimensions.items()
                              if d.hidden),
        "n_images": len(getattr(ws, "_images", []) or []),
        "n_charts": len(getattr(ws, "_charts", []) or []),
        "cells": cells,
        "n_border_only": border_only,
        "n_formula_no_cache": n_formula_no_cache,
        "degraded": degrade,
    }


def _col_idx(letters):
    n = 0
    for ch in letters.upper():
        n = n * 26 + ord(ch) - 64
    return n


def _calamine_values(path, min_mb, rules_fast=True):
    """大文件快路径：calamine 读缓存值（实测比 openpyxl data_only 快约一个数量级）。
    不可用或未达阈值返回 None，走常规双加载。"""
    if not rules_fast or os.path.getsize(path) < min_mb * 1024 * 1024:
        return None
    try:
        from python_calamine import CalamineWorkbook
    except ImportError:
        return None
    try:
        cw = CalamineWorkbook.from_path(path)
        out = {}
        for name in cw.sheet_names:
            vals = {}
            for ri, row in enumerate(cw.get_sheet_by_name(name).to_python(), start=1):
                for ci, v in enumerate(row, start=1):
                    if v is not None:
                        vals[(ri, ci)] = v
            out[name] = vals
        return out
    except Exception:
        return None


def extract_workbook(path: str, rules=None) -> dict:
    if not os.path.exists(path):
        raise SystemExit(f"文件不存在: {path}")
    # xlsx 是 zip 包（PK 头）；.xls/.csv/损坏文件给出可读错误而非裸 traceback
    with open(path, "rb") as f:
        if f.read(2) != b"PK":
            raise SystemExit(f"不是 xlsx 格式（.xls/.csv 需先经 LibreOffice 归一化，该路径尚未接入）: {path}")
    from .rules import get_rules
    rules = rules or get_rules()
    safe_path, sanitized = sanitize_xlsx(path)
    wb = openpyxl.load_workbook(safe_path, data_only=False)
    values_degraded = False
    cal = _calamine_values(safe_path, rules.get("extract", {}).get("fast_values_min_mb", 4))
    if cal is not None:
        wbv = None  # 值走 calamine
    else:
        try:
            wbv = openpyxl.load_workbook(safe_path, data_only=True)
        except Exception:
            wbv = wb  # 缓存值加载失败时降级：问答数值精度受损但几何不受影响
            values_degraded = True
    sheets = []
    for name in wb.sheetnames:
        ws = wb[name]
        cal_values = cal.get(name) if cal is not None else None
        wsv = None
        if cal_values is None:
            if wbv is None:
                # 走了 calamine 总路径但该页不在结果里（图表页等）：必须显式降级
                wsv = ws
                values_degraded = True
            else:
                try:
                    wsv = wbv[name]
                except Exception:
                    wsv = ws
                    values_degraded = True
        s = extract_sheet(ws, wsv, cal_values)
        s["visible"] = wb[name].sheet_state == "visible"
        s["fast_values"] = cal_values is not None
        # 机打报表（openpyxl 等生成）公式无缓存值：几何不受影响，但数值层不可用，必须显式降级
        if s.get("n_formula_no_cache"):
            values_degraded = True
        sheets.append(s)
    return {"file": path, "sanitized": sanitized,
            "values_degraded": values_degraded, "sheets": sheets}
