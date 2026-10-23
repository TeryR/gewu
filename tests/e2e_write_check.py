# -*- coding: utf-8 -*-
"""端到端改数验证（真实 LLM + 真实 COM 写入）：
起服务 → 冻结 → 自然语言写计划 → 确认 diff → 执行写入 → 文件落盘核对 → 撤销 → 恢复核对。"""
import json
import os
import shutil
import subprocess
import sys
import time
import urllib.request

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)

SB = os.path.join("tests", "sandbox", "e2e_write.xlsx")
shutil.copy2(os.path.join("samples", "示例_集团费用归并.xlsx"), SB)
PORT = 8799
BASE = f"http://127.0.0.1:{PORT}"


def call(path, payload=None):
    if payload is None:
        r = urllib.request.urlopen(BASE + path)
        return json.loads(r.read().decode("utf-8"))
    req = urllib.request.Request(BASE + path, data=json.dumps(payload).encode("utf-8"),
                                 headers={"Content-Type": "application/json"}, method="POST")
    return json.loads(urllib.request.urlopen(req).read().decode("utf-8"))


proc = subprocess.Popen([sys.executable, "serve.py", SB, str(PORT)],
                        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
try:
    for _ in range(60):
        try:
            if call("/api/boot").get("ready"):
                break
        except Exception:
            pass
        time.sleep(0.5)
    else:
        raise SystemExit("服务未就绪")
    print("① 服务就绪")

    fr = call("/api/freeze", {"edits": []})
    assert fr.get("ok"), fr
    print("② 冻结完成")

    t0 = time.time()
    p = call("/api/plan_write", {"mode": "nl", "text": "把销售部1月差旅费的金额改成12345"})
    dt = time.time() - t0
    print(f"③ 写计划（真实 LLM，{dt:.1f}s）: ok={p.get('ok')}")
    assert p.get("ok"), p
    op = p["op"]
    print(f"   定位 {op['sheet']}!r{op['row']}c{op['col']}  旧值→新值 "
          f"{p.get('diff', {}).get('old', op.get('old'))} → {p.get('diff', {}).get('new', op.get('value'))}")
    print(f"   校验: {'放行' if p['validation']['ok'] else p['validation']['rejections']}"
          f"  警告: {p['validation'].get('warnings') or '无'}")
    assert p["validation"]["ok"] and op["sheet"] == "销售部" and op["row"] == 3, (op, p["validation"])

    e = call("/api/execute_write", {"op": op})
    print(f"④ 执行: ok={e.get('ok')} engine={e.get('engine')} 备份={os.path.basename(e.get('backup', ''))}")
    assert e.get("ok"), e

    import openpyxl
    v = openpyxl.load_workbook(SB, data_only=True)["销售部"].cell(row=3, column=3).value
    print(f"⑤ 文件落盘核对: 销售部 r3c3 = {v}（期望 12345）")
    assert v == 12345, v

    u = call("/api/undo", {})
    print(f"⑥ 撤销: {u}")
    assert u.get("ok") or u.get("restored"), u
    v2 = openpyxl.load_workbook(SB, data_only=True)["销售部"].cell(row=3, column=3).value
    print(f"⑦ 恢复核对: r3c3 = {v2}（期望 48222.78）")
    assert abs(v2 - 48222.78) < 0.01, v2

    print("\n端到端改数验证：全部通过 ✓（计划→写入→落盘→撤销→恢复）")
finally:
    proc.terminate()
    try:
        proc.wait(timeout=5)
    except Exception:
        proc.kill()
