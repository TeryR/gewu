# -*- coding: utf-8 -*-
"""框微调端到端实测（真服务 + 真 COM）：被空行劈开的表 → API 冻结合并 →
state 反映 → 规则引擎问答含下半 → 写入原框外行 → 撤销恢复。"""
import json
import os
import shutil
import subprocess
import sys
import time
import urllib.parse
import urllib.request

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)

from openpyxl import Workbook

SB = os.path.join("tests", "sandbox",
                  f"e2e_geom_{time.strftime('%H%M%S')}.xlsx")  # 每次唯一：避免快启复用上次状态
wb = Workbook()
ws = wb.active
ws.title = "台账"
ws["A1"] = "台账"
ws["A2"] = "项目"; ws["B2"] = "数量"
for i, v in enumerate([10, 20, 30]):
    ws.cell(row=3 + i, column=1, value=f"甲{i + 1}")
    ws.cell(row=3 + i, column=2, value=v)
for i, v in enumerate([40, 50]):
    ws.cell(row=7 + i, column=1, value=f"乙{i + 1}")
    ws.cell(row=7 + i, column=2, value=v)
wb.save(SB)

PORT = 8798
BASE = f"http://127.0.0.1:{PORT}"


def call(path, payload=None, method=None):
    if payload is None:
        r = urllib.request.urlopen(BASE + path)
        return json.loads(r.read().decode("utf-8"))
    req = urllib.request.Request(BASE + path, data=json.dumps(payload).encode("utf-8"),
                                 headers={"Content-Type": "application/json"},
                                 method=method or "POST")
    return json.loads(urllib.request.urlopen(req).read().decode("utf-8"))


proc = subprocess.Popen([sys.executable, "serve.py", SB, str(PORT)],
                        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
try:
    for _ in range(60):
        try:
            if call("/api/boot").get("ready"):
                break
        except Exception:
            pass
        time.sleep(0.5)
    else:
        raise SystemExit("服务未就绪")
    print("① 服务就绪")

    q = urllib.parse.quote("台账")
    st = call(f"/api/state?sheet={q}")
    tables = [b for b in st["blocks"] if b["type"] == "table"]
    print(f"② 调整前：{len(tables)} 个表块 {[b['bbox'] for b in tables]}")
    assert len(tables) == 2, "前提：表应被空行劈成两块"
    b1 = min(tables, key=lambda b: b["bbox"][0])
    b2 = max(tables, key=lambda b: b["bbox"][0])

    r = call("/api/freeze", {"edits": [
        {"block_db_id": b2["db_id"], "user_type": "ignore"},
        {"block_db_id": b1["db_id"], "user_bbox": [1, 1, 8, 2]},
    ]})
    print(f"③ 冻结（忽略下半 + 上半扩全表）: applied={r['applied']} rejected={r['rejected']}")
    assert r["applied"] == 2 and not r["rejected"], r

    st2 = call(f"/api/state?sheet={q}")
    t2 = [b for b in st2["blocks"] if b["type"] == "table"]
    ign = [b for b in st2["blocks"] if b["type"] == "ignore"]
    print(f"④ 调整后：表×{len(t2)} bbox={t2[0]['bbox'] if t2 else None} "
          f"引擎原框={t2[0].get('engine_bbox') if t2 else None} 忽略×{len(ign)}")
    assert len(t2) == 1 and t2[0]["bbox"] == [1, 1, 8, 2] and len(ign) == 1

    a = call("/api/ask", {"question": "数量列合计", "sheet": "台账"}, "POST")
    print(f"⑤ 问答：{a['answer'].splitlines()[0]}（期望 150，含劈掉的下半）")
    assert "150" in a["answer"], a["answer"]

    p = call("/api/plan_write", {"mode": "nl", "text": "把乙2的数量改成77"})
    if not p.get("ok"):                    # LLM 单次输出变体触发回问 → 重试一次
        p = call("/api/plan_write", {"mode": "nl", "text": "把台账表里项目为乙2的数量改成77"})
    print(f"⑥ 写计划（真实 LLM）：ok={p.get('ok')} 定位={p.get('op', {}).get('sheet')}"
          f"!r{p.get('op', {}).get('row')}c{p.get('op', {}).get('col')}")
    assert p.get("ok") and p["op"]["row"] == 8, p
    e = call("/api/execute_write", {"op": p["op"]})
    assert e.get("ok"), e
    from openpyxl import load_workbook
    v = load_workbook(SB, data_only=True)["台账"]["B8"].value
    print(f"⑦ 写入落盘 B8={v}（r8 在引擎框外、人工框内——调整前这里会被拒）")
    assert v == 77

    u = call("/api/undo", {})
    v2 = load_workbook(SB, data_only=True)["台账"]["B8"].value
    print(f"⑧ 撤销恢复 B8={v2}")
    assert u.get("ok") and v2 == 50

    print("\n框微调端到端实测：全部通过 ✓")
finally:
    proc.terminate()
    try:
        proc.wait(timeout=5)
    except Exception:
        proc.kill()
