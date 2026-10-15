# -*- coding: utf-8 -*-
"""数据解读（op=analyze）：『有什么评价/哪个更好』不再被塞进 lookup 倒原始记录。

数值纪律：统计量（n/min/max/avg）由引擎计算，LLM 只引用；records 百分比按 % 格式化。"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import Workbook

from engine import llm
from engine.mirror import build_mirror, freeze
from engine.query import Mirror

SB = os.path.join("tests", "sandbox")
os.makedirs(SB, exist_ok=True)
PATH = os.path.join(SB, "analyze.xlsx")
DB = os.path.join(SB, "analyze.db")


def _build():
    wb = Workbook()
    ws = wb.active
    ws.title = "对比表"
    ws["A1"] = "样品"; ws["B1"] = "强力"; ws["C1"] = "变化率"
    for i, (s, f, r) in enumerate([("甲", 100, -0.10), ("乙", 120, -0.05),
                                   ("丙", 80, -0.30), ("丁", 95, -0.12)]):
        ws.cell(row=2 + i, column=1, value=s)
        ws.cell(row=2 + i, column=2, value=f)
        c = ws.cell(row=2 + i, column=3, value=r)
        c.number_format = "0%"
    wb.save(PATH)


print("== 构建：4 行对比表（含百分比列） ==")
_build()
if os.path.exists(DB):
    os.remove(DB)
wid = build_mirror(PATH, DB)
freeze(wid, DB)
m = Mirror(DB, wid)

print("\n== 链路1：LLM 判 analyze → analyze_block 收到引擎统计量 ==")
llm.chat = lambda messages, **kw: '{"op":"analyze","block_hint":"对比表"}'
seen = {}

def fake_analyze(question, ctx, history=None):
    seen["q"] = question
    seen["ctx"] = ctx
    return "乙样强力最高且降幅最小，整体工艺稳定。"

llm.analyze_block = fake_analyze
ans = m.answer("你对这些数据有什么评价吗？", llm=llm.query_intent)
assert ans == "乙样强力最高且降幅最小，整体工艺稳定。"
ctx = seen["ctx"]
st = ctx["列统计"]["强力"]
assert st["n"] == 4 and st["max"] == 120 and st["min"] == 80, st
assert abs(float(st["avg"]) - 98.75) < 1e-9, st
r = ctx["列统计"]["变化率"]
assert "%" in str(r["min"]) and "%" in str(r["avg"]), r   # 百分比列按 % 展示
assert ctx["records"][0]["变化率"] == "-10.00%" or "-10" in str(ctx["records"][0]["变化率"])
print(f"  强力统计={st}  变化率统计min={r['min']}")
print(f"  records样例={ctx['records'][0]}")

print("\n== 链路2：LLM 挂了 → 引擎统计兜底 ==")
def boom(question, ctx, history=None):
    raise SystemExit("API 超时")

llm.analyze_block = boom
ans = m.answer("评价一下这表", llm=llm.query_intent)
assert "[对比表]" in ans and "均值" in ans and "98.75" in ans
print("  ✓ 兜底输出统计表")

print("\n== 链路3：无 LLM 规则版关键词路由 ==")
ans = m.answer("分析评价一下这个表", llm=None)
assert "引擎统计" in ans
print("  ✓ 规则版走统计兜底")

print("\n== 链路4：sheet_hint 对 analyze 也生效 ==")
llm.chat = lambda messages, **kw: '{"op":"analyze"}'
llm.analyze_block = fake_analyze
ans = m.answer("有什么评价吗", llm=llm.query_intent, sheet_hint="对比表")
assert ans == "乙样强力最高且降幅最小，整体工艺稳定。"
print("  ✓ 无 block_hint 时落当前页")

print("\n== 链路5：全簿设计评价（block_hint=*） ==")
llm.chat = lambda messages, **kw: '{"op":"analyze","block_hint":"*"}'
seen5 = {}
def fake_wb_analyze(question, ctx, history=None):
    seen5['ctx'] = ctx
    return '该文件按规格×对象展开变体表，结构重复但口径一致。'
llm.analyze_block = fake_wb_analyze
ans = m.answer("你觉得设计合理吗", llm=llm.query_intent)
assert "变体表" in ans
ctx5 = seen5['ctx']
assert ctx5.get('workbook') is True and ctx5['n_sheets'] >= 1
assert len(ctx5['表结构画像']) >= 1 and '近亲sheet分组' in ctx5
print(f"  ✓ 全簿上下文: {ctx5['n_sheets']} sheets, 画像{len(ctx5['表结构画像'])}表")

print("\n== 链路6：无目标评价默认全簿（无范围无当前页） ==")
llm.chat = lambda messages, **kw: '{"op":"analyze"}'
llm.analyze_block = fake_wb_analyze
ans = m.answer("评价一下整体设计", llm=llm.query_intent)
assert "变体表" in ans
print('  ✓ 无目标+无范围 → block_hint=*')

print("\n全部通过 ✓")
