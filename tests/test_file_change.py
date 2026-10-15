# -*- coding: utf-8 -*-
"""文件变化处理：冻结后文件被改（自己写数 / 外部编辑）的接管策略。

核心原则：冻结保护的是【结构判定】——
- 纯数值变化 → 结构不动 → 冻结自动保持，数值随重建刷新
- 结构变化（块增/删/列变）→ 降级待复核；未变块的人工标记自动迁移
- 自己写数 → apply_ops_to_mirror 滚动 hash + 只更新受影响记录
- 写前护栏 → 文件被外部改过而未重新解析 → 拒绝写入"""
import json
import os
import sqlite3
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import Workbook, load_workbook

from engine import writer
from engine.mirror import build_mirror, freeze, apply_ops_to_mirror
from engine.query import Mirror

SB = os.path.join("tests", "sandbox")
os.makedirs(SB, exist_ok=True)
PATH = os.path.join(SB, "file_change.xlsx")
DB = os.path.join(SB, "file_change.db")


def _build():
    wb = Workbook()
    ws = wb.active
    ws.title = "台账"
    ws["A1"] = "项目"; ws["B1"] = "月份"; ws["C1"] = "数量"
    for i, row in enumerate([("甲", "1月", 10), ("乙", "1月", 20), ("丙", "1月", 30)]):
        for j, v in enumerate(row):
            ws.cell(row=2 + i, column=1 + j, value=v)
    wb.save(PATH)


def _status(db, wid):
    con = sqlite3.connect(db)
    st, ch = con.execute("SELECT status, change_json FROM workbooks WHERE id=?",
                         (wid,)).fetchone()
    con.close()
    return st, (json.loads(ch) if ch else None)


print("== 基线：构建 + 冻结 ==")
_build()
if os.path.exists(DB):
    os.remove(DB)
wid = build_mirror(PATH, DB)
freeze(wid, DB)
m = Mirror(DB, wid)
b = m.blocks("台账", tables_only=True)[0]
assert m.aggregate(b, "数量", "sum")["value"] == 60

print("\n== 场景1：外部纯数值修改 → 冻结保持 + 数值刷新 ==")
wbk = load_workbook(PATH)
wbk["台账"]["C2"] = 15        # 10 → 15
wbk.save(PATH)
wid2 = build_mirror(PATH, DB)
assert wid2 == wid, "同路径接管应复用 wid"
st, ch = _status(DB, wid)
print(f"  status={st} change={ch}")
assert st == "frozen" and ch and ch["type"] == "value-only"
m2 = Mirror(DB, wid)
b2 = m2.blocks("台账", tables_only=True)[0]
assert m2.aggregate(b2, "数量", "sum")["value"] == 65, "数值应随重建刷新"

print("\n== 场景2：外部结构修改 → 降级待复核 + 标记迁移 ==")
# 先人工标记：整个表设只读（模拟用户确认时的手动修正）
con = sqlite3.connect(DB)
tbl = con.execute("SELECT id FROM blocks WHERE workbook_id=? AND sheet_name='台账'",
                  (wid,)).fetchone()[0]
con.execute("UPDATE blocks SET readonly=1 WHERE id=?", (tbl,))
con.commit(); con.close()
wbk = load_workbook(PATH)
ws2 = wbk.create_sheet("新表")
ws2["A1"] = "x"; ws2["B1"] = "y"; ws2["A2"] = 1; ws2["B2"] = 2; ws2["A3"] = 3; ws2["B3"] = 4
wbk.save(PATH)
wid3 = build_mirror(PATH, DB)
assert wid3 == wid
st, ch = _status(DB, wid)
print(f"  status={st} change={json.dumps(ch, ensure_ascii=False)[:150]}")
assert st == "parsed", "结构变化必须降级待复核"
assert ch and ch["type"] == "structure" and len(ch["added"]) >= 1
con = sqlite3.connect(DB)
ro = con.execute("SELECT readonly FROM blocks WHERE workbook_id=? AND sheet_name='台账'",
                 (wid,)).fetchone()
con.close()
assert ro and ro[0] == 1, "未变块的人工标记（只读）应迁移保留"
m3 = Mirror(DB, wid)
gl_before = len(m3.glossary())
print(f"  词典条目随接管保留: {gl_before}（同 wid 天然保留）")

print("\n== 场景3：复核后重新冻结（change_json 清空→横幅消失） ==")
freeze(wid, DB)
st, ch = _status(DB, wid)
assert st == "frozen" and ch is None, f"重冻后横幅未清: {ch}"
print("  ✓ 重冻后 change_json=None")

print("\n== 场景4：apply_ops_to_mirror → hash 滚动，重开即同 hash 续用 ==")
# 直接调引擎同步（不经 COM）：把 C2 15 改成 25
op = {"sheet": "台账", "row": 2, "col": 3, "value": 25}
apply_ops_to_mirror(wid, DB, PATH, [op])
m4 = Mirror(DB, wid)
b4 = m4.blocks("台账", tables_only=True)[0]
rec2 = next(r for r in m4._records(b4) if r["_r"] == 2)
assert rec2["数量"] == 25, "镜像记录应更新"
# 但注意：文件本身没真改（模拟里没写盘）——用真实写盘来验 hash：
wbk = load_workbook(PATH); wbk["台账"]["C2"] = 25; wbk.save(PATH)
apply_ops_to_mirror(wid, DB, PATH, [op])   # 写盘后再滚 hash
wid5 = build_mirror(PATH, DB)
con = sqlite3.connect(DB)
n = con.execute("SELECT COUNT(*) FROM workbooks WHERE path=?", (PATH,)).fetchone()[0]
con.close()
assert wid5 == wid and n == 1, "hash 已滚动 → 重开走同 hash 分支，不产生新工作簿"

print("\n== 场景5：写前护栏 → 外部改动未重解析 → 拒绝写入 ==")
wbk = load_workbook(PATH); wbk["台账"]["C3"] = 999; wbk.save(PATH)  # 外部改，不 build_mirror
try:
    writer.execute([{"sheet": "台账", "row": 4, "col": 3, "value": 1}], PATH, DB, wid)
    raise AssertionError("应被 hash 护栏拒绝")
except writer.Rejection as e:
    print(f"  ✓ 拒绝: {str(e)[:46]}…")
    assert "外部修改" in str(e)

# 撤销同样被护栏拦：外部改动未重解析时，把旧值写回去可能砸掉别人的改动
from engine.writer import undo_last_batch
u = undo_last_batch(DB, PATH, workbook_id=wid)
assert not u["ok"] and "外部修改" in u["reason"], u
print(f"  ✓ undo 护栏: {u['reason'][:36]}…")

# 还原：重解析（value-only）→ 冻结保持 → 写入放行
wid6 = build_mirror(PATH, DB)
st, ch = _status(DB, wid)
assert st == "frozen" and ch["type"] == "value-only"

print("\n全部通过 ✓")
