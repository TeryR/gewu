# -*- coding: utf-8 -*-
"""多级表头叶子行不得被孤点剥离误伤（多级表头对比汇总表模式）。

背景：四级表头（r2 组 → r3 子组 → r4 叶子标签），叶子列里"处理记录/复合/评价"
等列整列无数据 → r4 的表头格是该列唯一内容格 → 旧孤点规则把它们剥成黄色批注块，
几何上嵌在蓝表里（用户看到"蓝框套黄框"）。
判别信号：真批注（小结/更新日期）同一家不会复现；表头词汇在同一行横向重复
（header_rep_frac ≥ 0.5 的行整行不剥）。
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import Workbook

from engine.rules import get_rules
from engine.extract import extract_workbook
from engine.segment import segment_sheet
from engine.normalize import sanitize_xlsx

SB = os.path.join("tests", "sandbox")
os.makedirs(SB, exist_ok=True)
PATH = os.path.join(SB, "header_carve.xlsx")

HEADER_LEAF = (["基准", "处理", "处理记录", "复合", "评价", "基准", "处理",
                "平处变量", "处理记录", "复合", "平复变量", "评价", "基准", "处理",
                "绒处变量", "评价", "处理记录", "复合"])


def _build():
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    ws.merge_cells("A1:R1")
    ws["A1"] = "面料工艺性能对比"
    ws.merge_cells("A2:E2"); ws["A2"] = "A系列"
    ws.merge_cells("F2:J2"); ws["F2"] = "B系列"
    ws.merge_cells("K2:R2"); ws["K2"] = "分析"
    ws["A3"] = "序号"; ws["B3"] = "项目"; ws["C3"] = "基布样"; ws["D3"] = "处理液及规格"
    for i, lab in enumerate(HEADER_LEAF):
        ws.cell(row=4, column=5 + i, value=lab)
    data = [("厚度(mm)", 0.74, None, None, None, None, None, 0.86, None, None,
             None, None, 0.65, None, None, None, None, None),
            ("克重(g/㎡)", 250, None, None, None, None, None, 255, None, None,
             None, None, 180, None, None, None, None, None),
            ("伸长率(%)", 97, 73, None, None, None, None, 66, None, None,
             None, None, 79, None, None, None, None, None)]
    for j, row in enumerate(data):
        ws.cell(row=5 + j, column=1, value=j + 1)
        ws.cell(row=5 + j, column=2, value=row[0])
        ws.cell(row=5 + j, column=3, value="基布样")
        ws.cell(row=5 + j, column=4, value="TR-4020-130")
        for i, v in enumerate(row[1:]):
            if v is not None:
                ws.cell(row=5 + j, column=5 + i, value=v)
    # 个别"处理记录/复合"叶子列带数据 → 表头树应能命名它们（c7=处理记录 c14=复合）
    ws.cell(row=5, column=7, value="处1#")
    ws.cell(row=5, column=14, value="复1#")
    # 中部真批注行：小结 + 长备注，同一家只出现一次 → 必须被剥离
    ws.cell(row=9, column=3, value="小结")
    ws.cell(row=9, column=18, value="该批次原料吸湿偏高;处理后强力损失明显")
    wb.save(PATH)


def _blocks():
    rules = get_rules()
    cache, _ = sanitize_xlsx(PATH)
    wb = extract_workbook(cache, rules=rules)
    return segment_sheet(wb["sheets"][0], rules=rules)["blocks"]


print("== 构建：四级表头 + 无数据叶子列 + 中部小结行 ==")
_build()
blocks = _blocks()
notes = [(b["bbox"], b) for b in blocks if b["type"] == "note"]
tables = [b for b in blocks if b["type"] == "table"]
print(f"  区块: 表×{len(tables)} 注×{len(notes)} {[b[0] for b in notes]}")

# 1) 表头行 r4 不得被剥成批注
bad = [bb for bb, _ in notes if bb[0] == 4]
print(f"  1) 表头行r4被误剥: {bad if bad else '无 ✓'}")
assert not bad, f"表头叶子行被剥离为批注: {bad}"

# 2) 真小结行 r9 必须被剥
ok_note = any(bb[0] == 9 for bb, _ in notes)
print(f"  2) 小结行r9剥离: {'✓' if ok_note else '✗ 未剥离'}")
assert ok_note, "真批注（小结行）未被剥离"

# 3) 带数据的叶子列（处理记录/复合）保留在表内 → 表头树能命名它们
from engine.headers import build_header_tree
rules = get_rules()
cache, _ = sanitize_xlsx(PATH)
sh = extract_workbook(cache, rules=rules)["sheets"][0]
seg = segment_sheet(sh, rules=rules)
tab = next(b for b in seg["blocks"] if b["type"] == "table")
h = build_header_tree(sh, tab, rules)
joined = "".join(str(n.get("path")) + str(n.get("label") or "") for n in h.get("col_tree") or [])
kept = all(w in joined for w in ("处理记录", "复合"))
print(f"  3) 表头树含被保护的叶子词: {'✓' if kept else '✗ ' + joined[:160]}")
assert kept

print("\n全部通过 ✓")
