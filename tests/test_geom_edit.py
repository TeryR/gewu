# -*- coding: utf-8 -*-
"""框微调（user_bbox）引擎级回归。

场景：一张表被内部全空行劈成两块（引擎真实会发生的切错）→
人工调整范围合并 → 表头/记录在新范围重建 → 重建/接管幸存 →
写入原引擎框外行 → 重叠/非法范围拒绝。"""
import json
import os
import shutil
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import Workbook

from engine.extract import extract_workbook
from engine.mirror import build_mirror, freeze
from engine.query import Mirror
from engine import writer

SB = os.path.join("tests", "sandbox")
os.makedirs(SB, exist_ok=True)
PATH = os.path.join(SB, "geom_split.xlsx")
DB = os.path.join(SB, "geom.db")

NUMS_TOP = [10, 20, 30]        # r3-r5
NUMS_BOT = [40, 50]            # r7-r8（中间 r6 整行空 → 劈开）
TOTAL = sum(NUMS_TOP) + sum(NUMS_BOT)


def _build():
    wb = Workbook()
    ws = wb.active
    ws.title = "台账"
    ws["A1"] = "台账"
    ws["A2"] = "项目"; ws["B2"] = "数量"
    for i, v in enumerate(NUMS_TOP):
        ws.cell(row=3 + i, column=1, value=f"甲{i + 1}")
        ws.cell(row=3 + i, column=2, value=v)
    # r6 整行空 → 空带切断
    for i, v in enumerate(NUMS_BOT):
        ws.cell(row=7 + i, column=1, value=f"乙{i + 1}")
        ws.cell(row=7 + i, column=2, value=v)
    wb.save(PATH)


def _blocks(m):
    return {(b["block_id"]): b for b in m.blocks("台账", tables_only=True)}


print("== 前置：引擎把表劈成两块 ==")
_build()
if os.path.exists(DB):
    os.remove(DB)
wid = build_mirror(PATH, DB)
m = Mirror(DB, wid)
bs = _blocks(m)
print(f"  表块: {[(b['block_id'], b['r1'], b['r2']) for b in bs.values()]}")
assert len(bs) == 2, f"测试前提失败：应劈成两块，实际 {len(bs)}"
b1 = min(bs.values(), key=lambda b: b["r1"])
b2 = max(bs.values(), key=lambda b: b["r1"])
db1, db2 = b1["id"], b2["id"]
agg = m.aggregate(b1, "数量", "sum")
assert agg["value"] == sum(NUMS_TOP), agg
print(f"  上半块合计={agg['value']}（只有上半，被劈开的代价）")

print("\n== ① 冻结时合并：忽略下半块 + 上半块范围扩到全表 ==")
r = freeze(wid, DB, edits=[
    {"block_db_id": db2, "user_type": "ignore"},
    {"block_db_id": db1, "user_bbox": [1, 1, 8, 2]},
])
print(f"  freeze: applied={r['applied']} rejected={r['rejected']}")
assert r["applied"] == 2 and not r["rejected"], r
m2 = Mirror(DB, wid)
bs2 = _blocks(m2)
assert len(bs2) == 1, f"合并后应只剩一个表块，实际 {len(bs2)}"
bb = next(iter(bs2.values()))
assert [bb["r1"], bb["c1"], bb["r2"], bb["c2"]] == [1, 1, 8, 2], dict(bb)
agg2 = m2.aggregate(bb, "数量", "sum")
print(f"  合并后合计={agg2['value']}（期望 {TOTAL}，含劈掉的下半）")
assert agg2["value"] == TOTAL, agg2
# 表头树在新范围重建：列路径仍然可用
recs = m2._records(bb)
assert any(r.get("项目") == "乙2" for r in recs), "下半记录应进镜像"
print(f"  记录数={len(recs)}（含乙1/乙2）")

print("\n== ② 重建幸存（3 次幂等）==")
counts = []
for _ in range(3):
    wid2 = build_mirror(PATH, DB)
    assert wid2 == wid
    mm = Mirror(DB, wid2)
    bbs = _blocks(mm)
    assert len(bbs) == 1, bbs
    b_ = next(iter(bbs.values()))
    counts.append(m.aggregate(b_, "数量", "sum")["value"])
    mm.close()
print(f"  3 次合计: {counts}")
assert counts == [TOTAL] * 3

print("\n== ③ 接管幸存（外部改数值 → 冻结保持 + 几何保持）==")
from openpyxl import load_workbook
w = load_workbook(PATH)
w["台账"]["B3"] = 100          # 10 → 100
w.save(PATH)
wid3 = build_mirror(PATH, DB)
assert wid3 == wid, "同路径应接管复用 wid"
m3 = Mirror(DB, wid3)
import sqlite3
st = sqlite3.connect(DB).execute(
    "SELECT status FROM workbooks WHERE id=?", (wid3,)).fetchone()[0]
assert st == "frozen", f"纯数值变化冻结应保持：{st}"
bbs3 = _blocks(m3)
b3_ = next(iter(bbs3.values()))
assert [b3_["r1"], b3_["r2"]] == [1, 8], "人工几何应随接管幸存"
agg3 = m3.aggregate(b3_, "数量", "sum")
print(f"  接管后合计={agg3['value']}（期望 {TOTAL + 90}）")
assert agg3["value"] == TOTAL + 90, agg3

print("\n== ④ 写入原引擎框外行（合并前会被'无区块'拒绝）==")
op = {"sheet": "台账", "row": 8, "col": 2, "value": 999}   # r8 原属下半块
v = writer.validate_plan([op], PATH, DB, wid)
print(f"  校验: {'放行 ✓' if v['ok'] else '✗' + str(v['rejections'])}")
assert v["ok"], v["rejections"]
res = writer.execute([op], PATH, DB, wid)
assert res["applied"] == 1, res
pc = writer.post_check([op], PATH, DB, wid)
assert pc["ok"], pc["problems"]
wv = load_workbook(PATH, data_only=True)["台账"]["B8"].value
print(f"  写入落盘 B8={wv}，写后校验通过")
u = writer.undo_last_batch(DB, PATH, workbook_id=wid)
assert u.get("ok"), u
wv2 = load_workbook(PATH, data_only=True)["台账"]["B8"].value
assert wv2 == 50, f"撤销应恢复 50，实际 {wv2}"
print(f"  撤销恢复 B8={wv2} ✓")

print("\n== ⑤ 重叠拒绝（不先忽略就扩过去）==")
for mm in (m, m2, m3):
    try:
        mm.close()
    except Exception:
        pass
if os.path.exists(DB):
    os.remove(DB)
wid = build_mirror(PATH, DB)     # 撤销已恢复 B8；B3=100 保留
m = Mirror(DB, wid)
bs = _blocks(m)
db1b = min(bs.values(), key=lambda b: b["r1"])["id"]
r = freeze(wid, DB, edits=[{"block_db_id": db1b, "user_bbox": [1, 1, 8, 2]}])
print(f"  rejected: {r['rejected'][0]['reason'] if r['rejected'] else '无'}")
assert r["rejected"] and "重叠" in r["rejected"][0]["reason"], r
assert r["applied"] == 0

print("\n== ⑥ 非法/越界范围拒绝 ==")
for bad in ([5, 1, 3, 2], [0, 1, 3, 2], [1, 1, 999, 2], "x"):
    r = freeze(wid, DB, edits=[{"block_db_id": db1b, "user_bbox": bad}])
    assert r["rejected"], bad
print("  非法范围 ×4 全部被拒 ✓")

print("\n== ⑦ 文本块移动范围（不重建记录，只挪框）==")
wbk = load_workbook(PATH)
ws2 = wbn = wbk.create_sheet("说明页")
ws2["A1"] = "表级说明文本"
wbk.save(PATH)
wid4 = build_mirror(PATH, DB)    # 结构变化 → 接管
m4 = Mirror(DB, wid4)
notes = [b for b in m4.blocks("说明页") if b["type"] == "note"]
if notes:
    nr = freeze(wid4, DB, edits=[{"block_db_id": notes[0]["id"], "user_bbox": [1, 1, 2, 2]}])
    assert nr["applied"] == 1 and not nr["rejected"], nr
    m5 = Mirror(DB, wid4)
    n2 = [b for b in m5.blocks("说明页") if b["type"] == "note"][0]
    assert [n2["r1"], n2["c2"]] == [1, 2]
    print(f"  文本块范围 [1,1,2,2] ✓（只挪框，无记录重建）")
m4.close()

print("\n全部通过 ✓")
