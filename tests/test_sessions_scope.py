# -*- coding: utf-8 -*-
"""会话持久化 + 对话范围（scope）+ lookup 真实计数 + 跨文件历史搜索。"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import Workbook

from engine import llm
from engine.mirror import (build_mirror, chat_session_create, chat_session_list,
                           chat_session_load, chat_session_append,
                           chat_session_delete, chat_sessions_clear,
                           chat_history_search)
from engine.query import Mirror

SB = os.path.join("tests", "sandbox")
os.makedirs(SB, exist_ok=True)
PATH = os.path.join(SB, "scope_a.xlsx")
PATH_B = os.path.join(SB, "scope_b.xlsx")
DB = os.path.join(SB, "scope.db")


def _build(path, tag):
    wb = Workbook()
    for sn in ("台账A", "台账B"):
        ws = wb.create_sheet(sn)
        ws["A1"] = "项目"; ws["B1"] = "数量"
        for i in range(6):
            ws.cell(row=2 + i, column=1, value=f"{tag}-{sn}-{i}")
            ws.cell(row=2 + i, column=2, value=i + 1)
    wb.remove(wb["Sheet"])
    wb.save(path)


print("== 构建：双 sheet 台账（同列名）×两个文件 ==")
_build(PATH, "甲")
_build(PATH_B, "乙")
if os.path.exists(DB):
    os.remove(DB)
wid = build_mirror(PATH, DB)
wid_b = build_mirror(PATH_B, DB)

print("\n== 1) scope 隔离：范围=台账B 时问答只看 B ==")
llm.chat = lambda messages, **kw: '{"op":"sum","column":"数量","block_hint":"台账"}'
m = Mirror(DB, wid)
ans = m.answer("数量合计", llm=llm.query_intent, scope=["台账B"])
print(f"  范围B: {ans.splitlines()[0]}")
assert "[台账B]" in ans and "[台账A]" not in ans
ans_all = m.answer("数量合计", llm=llm.query_intent)
assert "[台账" in ans_all
print(f"  无范围: {ans_all.splitlines()[0]}")

print("\n== 2) 范围唯一 sheet 的『这个表格』指代 ==")
llm.chat = lambda messages, **kw: '{"op":"describe"}'
captured = {}
llm.describe_block = lambda q, c, history=None: (captured.update(sheet=c["sheet"]) or "ok")
m.answer("这个表格是做什么的", llm=llm.query_intent, scope=["台账B"])
assert captured["sheet"] == "台账B", captured
print(f"  ✓ describe 落在 {captured['sheet']}")

print("\n== 3) lookup 真实命中数（6 行全中，显示前 8 → 总数=6） ==")
llm.chat = lambda messages, **kw: '{"op":"lookup","column":"项目","block_hint":"台账B","filters":{}}'
ans = m.answer("列出所有", llm=llm.query_intent)
assert "命中 6 行" in ans and "最多显示" not in ans, ans[:120]
print(f"  ✓ {ans.splitlines()[0]}")

print("\n== 4) 会话 CRUD + 首问标题 + 持久化跨文件版本（接管同 wid） ==")
sid = chat_session_create(wid, DB, title="布料工艺的疑问", scope=["台账B"])
chat_session_append(sid, "user", "变量列怎么算的", DB)
chat_session_append(sid, "assistant", "=(浸胶-原样)/原样", DB)
sess = chat_session_list(wid, DB)
assert len(sess) == 1 and sess[0]["title"] == "布料工艺的疑问" and sess[0]["scope"] == ["台账B"]
msgs = chat_session_load(sid, DB)
assert len(msgs) == 2 and msgs[0]["role"] == "user"
# 文件结构变化 → 接管复用 wid → 会话仍在
from openpyxl import load_workbook
wbk = load_workbook(PATH)
wbn = wbn2 = wbk.create_sheet("新增表")
wbk.save(PATH)
wid2 = build_mirror(PATH, DB)
assert wid2 == wid
sess2 = chat_session_list(wid, DB)
assert len(sess2) == 1 and sess2[0]["title"] == "布料工艺的疑问"
print("  ✓ 结构变化接管后历史仍在")

print("\n== 5) 跨文件历史搜索（文件B 的会话能被文件A 侧搜到） ==")
sid_b = chat_session_create(wid_b, DB, title="乙文件的问题")
chat_session_append(sid_b, "user", "独家关键词蓝鲸物性", DB)
hits = chat_history_search("蓝鲸物性", DB)
assert hits and hits[0]["file"] == "scope_b.xlsx" and "独家关键词" in hits[0]["snippet"]
print(f"  ✓ 命中 [{hits[0]['file']}] {hits[0]['title']}")

print("\n== 6) 删除/清空 ==")
chat_session_delete(sid_b, DB)
assert not chat_history_search("蓝鲸物性", DB)
chat_sessions_clear(wid, DB)
assert not chat_session_list(wid, DB)
print("  ✓ 单删 + 清空生效")

print("\n== 7) 解耦验证：lookup 显示行数由 rules.json 驱动 ==")
from engine.rules import get_rules
rl = get_rules()
rl["query"]["lookup_show_rows"] = 3
llm.chat = lambda messages, **kw: '{"op":"lookup","column":"项目","block_hint":"台账B","filters":{}}'
ans = m.answer("列出所有", llm=llm.query_intent)
assert "显示前 3 行" in ans, ans[:100]
rl["query"]["lookup_show_rows"] = 8
print("  ✓ lookup_show_rows=3 → 只显示前 3 行（改配置不改代码）")

print("\n全部通过 ✓")
