# -*- coding: utf-8 -*-
"""业务规则层：人配置值域，引擎硬执行（block 拒写 / warn 警告），LLM 只知情。"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import Workbook

from engine import llm
from engine.mirror import build_mirror, freeze, add_biz_rule, list_biz_rules, del_biz_rule
from engine.query import Mirror
from engine.writer import validate_plan

SB = os.path.join("tests", "sandbox")
os.makedirs(SB, exist_ok=True)
P = os.path.join(SB, "biz_rule.xlsx")
D = os.path.join(SB, "biz_rule.db")

wb = Workbook(); ws = wb.active; ws.title = "台账"
ws["A1"] = "项目"; ws["B1"] = "数量"
for i, (k, v) in enumerate([("甲", 10), ("乙", 999), ("丙", 50)]):
    ws.cell(row=2 + i, column=1, value=k); ws.cell(row=2 + i, column=2, value=v)
wb.save(P)
if os.path.exists(D):
    os.remove(D)
wid = build_mirror(P, D)
freeze(wid, D)
OP = {"sheet": "台账", "row": 2, "col": 2}

print("== ① block 规则：越界拒写、界内放行 ==")
add_biz_rule(wid, "台账", "数量", rmin=0, rmax=100, severity="block", db=D)
v = validate_plan([{**OP, "value": 500}], P, D, wid)
assert not v["ok"] and "业务规则" in v["rejections"][0], v
print(f"  ✓ 拒: {v['rejections'][0][:52]}")
assert validate_plan([{**OP, "value": 80}], P, D, wid)["ok"]
print("  ✓ 界内 80 放行")

print("== ② warn 规则：越界仅警告 ==")
add_biz_rule(wid, "台账", "数量", rmin=0, rmax=100, severity="warn", db=D)
v = validate_plan([{**OP, "value": 500}], P, D, wid)
assert v["ok"] and any("业务规则" in w for w in v["warnings"]), v
print("  ✓ 放行+警告")

print("== ③ 全表体检：存量违规行（乙=999） ==")
m = Mirror(D, wid)
viol = m.rules_check()
assert len(viol) == 1 and viol[0]["value"] == 999 and viol[0]["row"] == 3, viol
print(f"  ✓ {viol[0]['sheet']} r{viol[0]['row']} {viol[0]['column']}={viol[0]['value']} {viol[0]['问题']}")

print("== ④ 规则入 LLM 写上下文（只知情不执行） ==")
ctx = llm.mirror_context(m)
assert any(rr["column"] == "数量" for rr in ctx.get("业务规则", [])), ctx.get("业务规则")
print(f"  ✓ mirror_context 含业务规则: {ctx['业务规则']}")

print("== ⑤ 覆盖式保存与删除 ==")
add_biz_rule(wid, "台账", "数量", rmin=0, rmax=50, severity="block", db=D)
rs = list_biz_rules(wid, db=D)
assert len(rs) == 1 and rs[0]["rmax"] == 50.0
v = validate_plan([{**OP, "value": 80}], P, D, wid)
assert not v["ok"], "收紧后 80 应被拒"
del_biz_rule(wid, "台账", "数量", D)
assert validate_plan([{**OP, "value": 9999}], P, D, wid)["ok"]
print("  ✓ upsert 收紧生效 / 删除后恢复自由")

m.close()
print("\n全部通过 ✓")
