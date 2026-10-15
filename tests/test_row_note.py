# -*- coding: utf-8 -*-
"""行内备注挂接：剥离的黄块挂回所属表行（records._note + row_notes + LLM 上下文）。

布局仿真实案例（总部·项目性质）：备注列夹在中间被剥开 → 表劈成左右两半，
挂接取列距最近的表（平手取左，行身份在左半）。
覆盖：正常挂接 / 小结整行不挂 / 表级说明词不挂 / 人工解除且重建幸存 /
上下文注入 / 解除后 _note 即时清除 / 数值路径不受影响。"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import Workbook

from engine.mirror import build_mirror, detach_note
from engine.query import Mirror

SB = os.path.join("tests", "sandbox")
os.makedirs(SB, exist_ok=True)
PATH = os.path.join(SB, "row_note.xlsx")
DB = os.path.join(SB, "row_note.db")


def _build():
    wb = Workbook()
    ws = wb.active
    ws.title = "台账"
    ws["A1"] = "项目"; ws["B1"] = "性质"; ws["D1"] = "签约额"; ws["E1"] = "回款"  # C 列无表头（真实案例即如此）
    rows = [("鹿鸣台", "独立核算", None, 100, 30),
            ("青枫苑", "独立核算", "2026年7月开始变更为纯渠道", 200, 60),
            ("锦云梁园", "合作", None, 300, 90)]
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            if v is not None:
                ws.cell(row=2 + i, column=1 + j, value=v)
    # 小结整行（该行无其他表内容）→ 剥离但不挂
    ws.cell(row=6, column=2, value="小结")
    ws.cell(row=6, column=5, value="本月整体达成率超预期")
    # 第二页：表级说明词 → 剥离但不挂
    ws2 = wb.create_sheet("说明页")
    ws2["A1"] = "项目"; ws2["C1"] = "数值"  # B 列无表头，仅一格表级说明
    ws2["A2"] = "甲"; ws2["B2"] = "本表数据截至2026年9月"; ws2["C2"] = 1
    ws2["A3"] = "乙"; ws2["C3"] = 2
    wb.save(PATH)


def _records(m, sheet="台账"):
    """全表记录按 (块,行) 摊平：左右并排表同行号不互相覆盖。"""
    out = {}
    for b in m.blocks(sheet, tables_only=True):
        for r in m._records(b):
            out[(b["block_id"], r["_r"])] = r
    return out


print("== 构建：中间备注列（表劈两半）+ 小结行 + 表级说明页 ==")
_build()
if os.path.exists(DB):
    os.remove(DB)
wid = build_mirror(PATH, DB)
m = Mirror(DB, wid)
notes = m.row_notes()
print(f"  挂接数={len(notes)} {[(n['row_no'], n['sheet_name'], n['text'][:14]) for n in notes]}")
assert len(notes) == 1, f"应只挂 1 条行内备注，实际 {notes}"
assert notes[0]["row_no"] == 3 and "纯渠道" in notes[0]["text"]
assert notes[0]["sheet_name"] == "台账"

# 挂到左半表（行身份 项目/性质 在左半）：含「项目」列的表，其 r3 记录带 _note
def _recs_by_block(m, sheet="台账"):
    out = []
    for b in m.blocks(sheet, tables_only=True):
        out.append((b, m._records(b)))
    return out

by_block = _recs_by_block(m)
left = next((rs for b, rs in by_block
             if any("项目" in k for r in rs for k in r)), None)
r3 = next((r for r in left if r["_r"] == 3), None)
assert r3 and r3.get("项目") == "青枫苑" and "纯渠道" in r3.get("_note", ""), \
    f"左半表 r3 应带 _note: {r3}"
right = next((rs for b, rs in by_block
              if any("签约额" in k for r in rs for k in r)), None)
assert all("_note" not in r for r in right), "右半表不应带备注"
assert not any("_note" in r for b, rs in by_block for r in rs if r["_r"] == 6), \
    "小结行不得注入"
print(f"  左半表 r3={json.dumps(r3, ensure_ascii=False)[:120]}")

ctx = m._llm_context()
remarks = ctx.get("row_remarks") or []
print(f"  LLM上下文 row_remarks n={len(remarks)}")
assert any("纯渠道" in r["备注"] for r in remarks)
assert ctx.get("tables"), "上下文结构完整"

print("\n== 人工解除：挂接消失、_note 即时清除、重建幸存 ==")
nb = notes[0]["note_block_id"]
r = detach_note(wid, "台账", nb, DB)
print(f"  detach → {r}")
assert r["detached"] == 1
m2 = Mirror(DB, wid)
assert m2.row_notes() == [], "解除后不应再有挂接"
assert not any("_note" in r for r in _records(m2).values()), "解除后 _note 应清除"

wid2 = build_mirror(PATH, DB)  # 同文件重建：人工解除必须幸存
assert wid2 == wid
m3 = Mirror(DB, wid2)
assert m3.row_notes() == [], "重建不得把人工解除的挂接加回来"
print("  重建后挂接仍为空 ✓")

print("\n== 数值路径不受 _note 影响 ==")
for mm in (m, m2, m3):
    try:
        mm.close()
    except Exception:
        pass
if os.path.exists(DB):
    os.remove(DB)
wid = build_mirror(PATH, DB)
m = Mirror(DB, wid)
b = m._pick_block_for_column("签约额")
agg = m.aggregate(b, "签约额", "sum")
print(f"  sum(签约额)={agg['value'] if agg else None}（期望600）")
assert agg and agg["value"] == 600

print("\n全部通过 ✓")
