# -*- coding: utf-8 -*-
"""对抗样本与死路径实测：可运行 ≠ 运行良好。
覆盖：空工作簿 / 仅空合并区 / 公式无缓存值（openpyxl 产物）/ 非 xlsx 格式 /
     死路径（lookup/filters/count）/ writer 批量与 WPS 引擎。"""
import os
import sys
import shutil

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from engine.report import analyze_file
from engine.extract import extract_workbook

SB = os.path.join("tests", "sandbox")
os.makedirs(SB, exist_ok=True)


def make(name, builder):
    p = os.path.join(SB, name)
    wb = openpyxl.Workbook()
    builder(wb)
    wb.save(p)
    return p


def case(title, fn):
    try:
        out = fn()
        print(f"  {out}  <- {title}")
    except SystemExit as e:
        print(f"  SystemExit({e})  <- {title}")
    except Exception as e:
        print(f"  ✗未处理异常 {type(e).__name__}: {str(e)[:70]}  <- {title}")


print("== A. 对抗样本 ==")

def b_empty(wb):
    pass  # 只有一个空默认页

def b_merged_empty(wb):
    ws = wb.active
    ws.merge_cells("B2:D4")

def b_no_cache(wb):
    ws = wb.active
    ws["A1"], ws["A2"], ws["A3"] = "项目", "甲", "乙"
    ws["B1"], ws["B2"], ws["B3"] = "数量", 1, 2
    ws["C1"], ws["C4"] = "合计", "=SUM(B2:B3)"  # openpyxl 写的公式没有缓存值
    ws["A5"], ws["B5"] = "备注：本表为系统生成", None

def b_only_numbers(wb):
    ws = wb.active
    for r in range(1, 6):
        for c in range(1, 6):
            ws.cell(row=r, column=c).value = r * c

for name, builder in [("empty.xlsx", b_empty), ("merged_empty.xlsx", b_merged_empty),
                      ("no_cache.xlsx", b_no_cache), ("only_numbers.xlsx", b_only_numbers)]:
    p = make(name, builder)
    case(name, lambda p=p: f"blocks={len(analyze_file(p)['sheets'][0]['blocks'])}, "
                           f"degraded={analyze_file(p).get('values_degraded')}")

print("\n== B. 非 xlsx / 坏文件 ==")
csv_path = os.path.join(SB, "fake.csv")
with open(csv_path, "w", encoding="utf-8") as f:
    f.write("a,b\n1,2\n")
case("fake.csv（内容其实是文本）", lambda: extract_workbook(csv_path))
case("不存在的文件", lambda: extract_workbook(os.path.join(SB, "nope.xlsx")))

print("\n== C. no_cache 深挖：勾稽与降级标记 ==")
p = make("no_cache2.xlsx", b_no_cache)
wb = extract_workbook(p)
f_cells = [(pos, i) for pos, i in wb["sheets"][0]["cells"].items() if i["f"]]
none_cv = [pos for pos, i in f_cells if i["cv"] is None]
print(f"  公式格 {len(f_cells)} 个，其中 cv=None {len(none_cv)} 个 → "
      f"values_degraded={wb.get('values_degraded')}（应为 True？）")
