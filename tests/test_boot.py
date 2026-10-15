# -*- coding: utf-8 -*-
"""启动双改：①快启（同文件+同引擎→直接载镜像，不重解析）②引擎签名驱动缓存失效。"""
import json
import os
import sqlite3
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import Workbook

from engine.mirror import build_mirror, engine_sig, try_fast_boot

SB = os.path.join("tests", "sandbox")
os.makedirs(SB, exist_ok=True)
P = os.path.join(SB, "fastboot.xlsx")
D = os.path.join(SB, "fastboot.db")

wb = Workbook(); ws = wb.active; ws.title = "t"
ws["A1"] = "k"; ws["B1"] = "v"; ws["A2"] = "甲"; ws["B2"] = 1
wb.save(P)
if os.path.exists(D):
    os.remove(D)

print("== ① 首次构建：写 engine_ver ==")
wid = build_mirror(P, D)
con = sqlite3.connect(D)
ver = con.execute("SELECT engine_ver FROM workbooks WHERE id=?", (wid,)).fetchone()[0]
con.close()
assert ver == engine_sig(), (ver, engine_sig())
print(f"  ✓ engine_ver={ver}")

print("== ② 快启命中：直接复用镜像（不重解析） ==")
# 哨兵：手改一个 block 的置信分——快启必须原样保留（证明没重建）
con = sqlite3.connect(D)
con.execute("UPDATE blocks SET score=0.42 WHERE workbook_id=?", (wid,))
con.commit(); con.close()
t0 = time.time()
fw = try_fast_boot(P, D)
dt = (time.time() - t0) * 1000
assert fw == wid, fw
con = sqlite3.connect(D)
score = con.execute("SELECT score FROM blocks WHERE workbook_id=? LIMIT 1", (wid,)).fetchone()[0]
con.close()
assert abs(score - 0.42) < 1e-9, f"快启竟重建了镜像（score={score}）"
print(f"  ✓ 命中快启 {dt:.0f}ms，哨兵保留（未重解析）")

print("== ③ 文件变化 → 快启未命中（走重建） ==")
wb2 = Workbook(); ws2 = wb2.active; ws2.title = "t2"
ws2["A1"] = "x"; ws2["B1"] = 1; ws2["A2"] = "y"; ws2["B2"] = 2
ws2["A3"] = "z"; ws2["B3"] = 3
wb2.save(P)
assert try_fast_boot(P, D) is None
print("  ✓ 内容变了快启拒绝")

print("== ④ 引擎签名变化 → 缓存失效（防旧结构镜像冒充新结果） ==")
build_mirror(P, D)  # 重建到新内容
con = sqlite3.connect(D)
con.execute("UPDATE workbooks SET engine_ver='deadbeef' WHERE id=?", (wid,))
con.commit(); con.close()
assert try_fast_boot(P, D) is None, "引擎版本不匹配必须拒绝快启"
print("  ✓ 版本不匹配拒绝快启")

print("== ⑤ engine_sig 对代码/规则变化敏感 ==")
sig1 = engine_sig()
rf = os.path.abspath("rules.json")
st = os.stat(rf)
os.utime(rf, (st.st_atime, st.st_mtime + 2))
sig2 = engine_sig()
os.utime(rf, (st.st_atime, st.st_mtime))  # 还原
assert sig1 != sig2, "rules.json 变化必须改变签名"
print("  ✓ rules.json 触碰即签名变（还原后下次构建刷新版本）")

print("\n全部通过 ✓")
